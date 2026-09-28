from io import BytesIO

import pymupdf
import pytest
from docx import Document

from slideai.core.errors import DomainError
from slideai.infrastructure.documents.parsers import LocatedText, parse_document


def test_scan_pdf_returns_no_extractable_text() -> None:
    document = pymupdf.open()
    document.new_page()
    pdf_bytes = document.tobytes()

    with pytest.raises(DomainError) as raised:
        parse_document("pdf", pdf_bytes)

    assert raised.value.code == "NO_EXTRACTABLE_TEXT"


def test_pdf_parser_preserves_page_number() -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Executive summary")

    located = parse_document("pdf", document.tobytes())

    assert located == [
        LocatedText(
            text="Executive summary",
            page_number=1,
            section_title=None,
            paragraph_index=1,
        )
    ]


def test_docx_parser_preserves_heading_and_paragraph_order() -> None:
    document = Document()
    document.add_heading("Market outlook", level=1)
    document.add_paragraph("Revenue is expected to grow.")
    stream = BytesIO()
    document.save(stream)

    located = parse_document("docx", stream.getvalue())

    assert located[0].section_title == "Market outlook"
    assert located[0].text == "Revenue is expected to grow."
    assert located[0].paragraph_index == 2


@pytest.mark.parametrize(
    ("extension", "content", "expected_section"),
    [
        ("md", "# Findings\n\nRevenue grew 12%.\n\n## Risks\n\nSupply remains tight.", "Findings"),
        ("txt", "First paragraph.\n\nSecond paragraph.", None),
    ],
)
def test_markdown_and_text_parser_return_located_paragraphs(
    extension: str, content: str, expected_section: str | None
) -> None:
    located = parse_document(extension, content.encode("utf-8"))

    assert located
    assert located[0].section_title == expected_section
    assert located[0].paragraph_index == 1


def test_chunker_respects_limit_overlap_and_location_boundaries() -> None:
    from slideai.infrastructure.documents.chunking import chunk_located_text

    chunks = chunk_located_text(
        [
            LocatedText("市场趋势。" * 500, page_number=1, section_title=None, paragraph_index=1),
            LocatedText("第二页内容。" * 20, page_number=2, section_title=None, paragraph_index=1),
        ],
        chunk_size=100,
        overlap=20,
    )

    assert len(chunks) > 2
    assert all(chunk.token_count <= 100 for chunk in chunks)
    assert {chunk.page_number for chunk in chunks if chunk.page_number == 2} == {2}
