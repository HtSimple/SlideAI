import re
from dataclasses import dataclass
from io import BytesIO
from typing import Protocol, cast

import pymupdf
from docx import Document

from slideai.core.errors import DomainError

_PARAGRAPH_BREAK = re.compile(r"\n\s*\n+")


@dataclass(frozen=True)
class LocatedText:
    text: str
    page_number: int | None
    section_title: str | None
    paragraph_index: int | None


class _PdfPage(Protocol):
    def get_text(self, option: str) -> str: ...


class _PdfDocument(Protocol):
    @property
    def page_count(self) -> int: ...

    def load_page(self, page_id: int) -> _PdfPage: ...


def parse_document(
    extension: str, content: bytes, *, max_chars: int = 2_000_000
) -> list[LocatedText]:
    normalized_extension = extension.lower().lstrip(".")
    if normalized_extension == "pdf":
        located = _parse_pdf(content)
    elif normalized_extension == "docx":
        located = _parse_docx(content)
    elif normalized_extension == "md":
        located = _parse_markdown(content)
    elif normalized_extension == "txt":
        located = _parse_text(content)
    else:
        raise DomainError("FILE_TYPE_MISMATCH", "The document format is not supported.")

    total_chars = sum(len(item.text) for item in located)
    if total_chars > max_chars:
        raise DomainError("FILE_TEXT_TOO_LARGE", "The extracted text exceeds the configured limit.")
    if not located or not any(item.text.strip() for item in located):
        raise DomainError("NO_EXTRACTABLE_TEXT", "No readable text was found in this file.")
    return located


def _parse_pdf(content: bytes) -> list[LocatedText]:
    try:
        with pymupdf.open(stream=content, filetype="pdf") as raw_document:
            document = cast(_PdfDocument, raw_document)
            located: list[LocatedText] = []
            for page_index in range(document.page_count):
                page = document.load_page(page_index)
                page_text = _normalize(page.get_text("text"))
                if page_text:
                    located.append(
                        LocatedText(
                            text=page_text,
                            page_number=page_index + 1,
                            section_title=None,
                            paragraph_index=1,
                        )
                    )
    except (pymupdf.FileDataError, ValueError) as error:
        raise DomainError("FILE_PARSE_FAILED", "The PDF file could not be read.") from error
    if not located:
        raise DomainError("NO_EXTRACTABLE_TEXT", "No readable text was found in this PDF.")
    return located


def _parse_docx(content: bytes) -> list[LocatedText]:
    try:
        document = Document(BytesIO(content))
    except Exception as error:
        raise DomainError("FILE_PARSE_FAILED", "The DOCX file could not be read.") from error

    section_title: str | None = None
    located: list[LocatedText] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        text = _normalize(paragraph.text)
        if not text:
            continue
        style_name = (paragraph.style.name or "") if paragraph.style else ""
        if style_name.lower().startswith("heading"):
            section_title = text
            continue
        located.append(
            LocatedText(
                text=text,
                page_number=None,
                section_title=section_title,
                paragraph_index=paragraph_index,
            )
        )
    if not located:
        raise DomainError("NO_EXTRACTABLE_TEXT", "No readable text was found in this DOCX.")
    return located


def _parse_markdown(content: bytes) -> list[LocatedText]:
    text = _decode_text(content)
    headings: list[tuple[int, str]] = []
    located: list[LocatedText] = []
    paragraph_index = 0
    pending: list[str] = []

    def flush() -> None:
        nonlocal paragraph_index
        body = _normalize("\n".join(pending))
        pending.clear()
        if body:
            paragraph_index += 1
            located.append(
                LocatedText(
                    text=body,
                    page_number=None,
                    section_title=" > ".join(title for _, title in headings) or None,
                    paragraph_index=paragraph_index,
                )
            )

    for line in text.splitlines():
        match = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if match:
            flush()
            level = len(match.group(1))
            headings = [(depth, title) for depth, title in headings if depth < level]
            headings.append((level, match.group(2).strip()))
        elif not line.strip():
            flush()
        else:
            pending.append(line)
    flush()
    if not located:
        raise DomainError(
            "NO_EXTRACTABLE_TEXT", "No readable text was found in this Markdown file."
        )
    return located


def _parse_text(content: bytes) -> list[LocatedText]:
    text = _decode_text(content)
    paragraphs = [
        _normalize(paragraph)
        for paragraph in _PARAGRAPH_BREAK.split(text.strip())
        if _normalize(paragraph)
    ]
    if not paragraphs:
        raise DomainError("NO_EXTRACTABLE_TEXT", "No readable text was found in this text file.")
    return [
        LocatedText(
            text=paragraph,
            page_number=None,
            section_title=None,
            paragraph_index=index,
        )
        for index, paragraph in enumerate(paragraphs, start=1)
    ]


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig", errors="strict")
    except UnicodeDecodeError as error:
        raise DomainError("FILE_ENCODING_INVALID", "Text files must use UTF-8 encoding.") from error


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()
